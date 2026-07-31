"""Fetch Service — parallel downloads, content-type detection, retries, cache."""

from __future__ import annotations

import io
import json
import os
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Any
from xml.etree import ElementTree

from app.faa.cache import DocumentCache
from app.faa.connectors.base import AcquisitionConnector
from app.faa.http_client import HttpClient
from app.faa.models import CandidateDocument, FetchedDocument, sha256_bytes, sha256_text, utc_now
from app.faa.web_enrichment import deepen_search_results, enrich_url, text_is_thin

_TAG_RE = re.compile(r"<[^>]+>")
_WS_RE = re.compile(r"\s+")


class FetchService:
    def __init__(
        self,
        cache: DocumentCache,
        *,
        live_fetch: bool = False,
        pdf_enabled: bool = True,
        client: HttpClient | None = None,
        connectors: dict[str, AcquisitionConnector] | None = None,
        max_workers: int = 6,
    ) -> None:
        self.cache = cache
        self.live_fetch = live_fetch
        self.pdf_enabled = pdf_enabled
        self.client = client or HttpClient()
        self.connectors = connectors or {}
        self.max_workers = max(1, min(max_workers, 12))
        self._fetch_samples: list[float] = []

    def fetch_many(self, candidates: list[CandidateDocument]) -> list[FetchedDocument]:
        if not candidates:
            return []
        if len(candidates) == 1 or self.max_workers == 1:
            return [self.fetch_one(c) for c in candidates]

        out: list[FetchedDocument | None] = [None] * len(candidates)
        with ThreadPoolExecutor(max_workers=self.max_workers) as pool:
            futs = {pool.submit(self.fetch_one, c): idx for idx, c in enumerate(candidates)}
            for fut in as_completed(futs):
                idx = futs[fut]
                try:
                    out[idx] = fut.result()
                except Exception as exc:
                    c = candidates[idx]
                    out[idx] = self._fail(c, f"parallel_fetch_error: {exc}")
        return [o for o in out if o is not None]

    def fetch_one(self, candidate: CandidateDocument) -> FetchedDocument:
        url = (candidate.url or "").strip()
        if not url:
            return self._fail(candidate, "missing_url")

        # Cache short-circuit by URL (non-search)
        if not url.startswith("search://"):
            skip, reason = self.cache.should_skip(url)
            if skip:
                row = self.cache.lookup(url) or {}
                return FetchedDocument(
                    candidate_id=candidate.candidate_id,
                    title=candidate.title,
                    url=url,
                    connector_id=candidate.connector_id,
                    document_type=candidate.document_type,
                    company=candidate.company,
                    symbol=candidate.symbol,
                    organisation=candidate.organisation,
                    published_at=candidate.published_at,
                    checksum=str(row.get("checksum") or ""),
                    etag=row.get("etag"),
                    last_modified=row.get("last_modified"),
                    skipped=True,
                    skip_reason=reason,
                    live_fetch=bool(row.get("live_fetch")),
                    metadata={"cache": row},
                )

        # Connector-specific fetch hook
        conn = self.connectors.get(candidate.connector_id)
        if conn is not None:
            try:
                custom = conn.fetch(candidate, self.client)
                if custom is not None:
                    ok, reason = conn.validate(custom)
                    if not ok:
                        conn.mark_failure(reason or "validate_failed")
                        custom.error = reason
                    else:
                        conn.mark_success(utc_now().isoformat())
                    return custom
            except Exception as exc:
                conn.mark_failure(str(exc))

        if url.startswith("search://"):
            return self._fetch_search(candidate)

        if self.live_fetch:
            return self._live_http_fetch(candidate)

        # Offline deterministic acquisition record (still processable; not seed corpus)
        text = self._offline_stub_text(candidate)
        checksum = sha256_text(text)
        return FetchedDocument(
            candidate_id=candidate.candidate_id,
            title=candidate.title,
            url=url,
            connector_id=candidate.connector_id,
            document_type=candidate.document_type,
            company=candidate.company,
            symbol=candidate.symbol,
            organisation=candidate.organisation,
            published_at=candidate.published_at or utc_now().date().isoformat(),
            content_type="text/plain",
            content_text=text,
            content_bytes_len=len(text.encode("utf-8")),
            checksum=checksum,
            live_fetch=False,
            metadata={
                "mode": "offline_acquisition_record",
                "note": "Set FAA_LIVE_FETCH=true for real HTTP/PDF downloads",
                "authority": (candidate.metadata or {}).get("authority"),
            },
        )

    def _live_http_fetch(self, candidate: CandidateDocument) -> FetchedDocument:
        url = candidate.url
        conn = self.connectors.get(candidate.connector_id)
        max_per_minute = conn.max_per_minute if conn else 30
        conditional = self.cache.conditional_headers(url)
        t0 = time.perf_counter()
        resp = self.client.get(
            url,
            connector_id=candidate.connector_id,
            max_per_minute=max_per_minute,
            conditional=conditional or None,
        )
        fetch_ms = (time.perf_counter() - t0) * 1000
        self._note_fetch(fetch_ms)

        if resp.error:
            if conn:
                conn.mark_failure(resp.error)
            return self._fail(candidate, f"live_fetch_error: {resp.error}", fetch_ms=fetch_ms, attempts=resp.attempts)

        etag = resp.header("etag")
        last_modified = resp.header("last-modified")
        if resp.status_code == 304:
            skip, reason = self.cache.should_skip(url, etag=etag, not_modified=True)
            row = self.cache.lookup(url) or {}
            if conn:
                conn.mark_success(utc_now().isoformat())
            return FetchedDocument(
                candidate_id=candidate.candidate_id,
                title=candidate.title,
                url=url,
                connector_id=candidate.connector_id,
                document_type=candidate.document_type,
                company=candidate.company,
                symbol=candidate.symbol,
                organisation=candidate.organisation,
                checksum=str(row.get("checksum") or ""),
                etag=etag or row.get("etag"),
                last_modified=last_modified or row.get("last_modified"),
                skipped=True,
                skip_reason=reason or "http_304_not_modified",
                live_fetch=True,
                fetch_ms=fetch_ms,
                attempts=resp.attempts,
                metadata={"http_status": 304},
            )

        raw = resp.content or b""
        checksum = sha256_bytes(raw)
        if self.cache.has_checksum(checksum):
            if conn:
                conn.mark_success(utc_now().isoformat())
            return FetchedDocument(
                candidate_id=candidate.candidate_id,
                title=candidate.title,
                url=url,
                connector_id=candidate.connector_id,
                document_type=candidate.document_type,
                company=candidate.company,
                symbol=candidate.symbol,
                organisation=candidate.organisation,
                checksum=checksum,
                etag=etag,
                last_modified=last_modified,
                skipped=True,
                skip_reason="checksum_exists",
                live_fetch=True,
                content_bytes_len=len(raw),
                fetch_ms=fetch_ms,
                attempts=resp.attempts,
                metadata={"http_status": resp.status_code},
            )

        ctype = (resp.header("content-type") or "").lower()
        detected = self._detect_content_type(url, ctype, raw)
        text = self._extract_text(detected, raw, url)
        enrich_meta: dict[str, Any] = {}
        # Strategic enrichment: thin / failed extract → Firecrawl → Browserbase
        min_chars = 800 if detected == "text/html" else 400
        if resp.status_code >= 400 or text_is_thin(text, min_chars=min_chars):
            page = enrich_url(self.client, str(resp.url or url))
            if page and page.get("markdown"):
                text = page["markdown"]
                detected = "text/markdown"
                enrich_meta = {
                    "enriched_by": page.get("source"),
                    "enrichment_format": page.get("format") or "markdown",
                }
        if resp.status_code >= 400 or not text.strip():
            text = text.strip() or (
                f"Live fetch HTTP {resp.status_code} from {url}. "
                f"Title: {candidate.title}. "
                "Insufficient extractable text; retained for provenance."
            )

        doc = FetchedDocument(
            candidate_id=candidate.candidate_id,
            title=candidate.title,
            url=str(resp.url or url),
            connector_id=candidate.connector_id,
            document_type=candidate.document_type,
            company=candidate.company,
            symbol=candidate.symbol,
            organisation=candidate.organisation,
            published_at=candidate.published_at or utc_now().date().isoformat(),
            content_type=detected,
            content_text=text[:200_000],
            content_bytes_len=len(raw),
            checksum=checksum,
            etag=etag,
            last_modified=last_modified,
            live_fetch=True,
            fetch_ms=fetch_ms,
            attempts=resp.attempts,
            metadata={
                "http_status": resp.status_code,
                "content_type": ctype,
                "detected_type": detected,
                "authority": (candidate.metadata or {}).get("authority"),
                **enrich_meta,
            },
        )
        if conn:
            ok, reason = conn.validate(doc)
            if ok:
                conn.mark_success(utc_now().isoformat())
            else:
                conn.mark_failure(reason or "validate_failed")
                doc.error = reason
        return doc

    def _detect_content_type(self, url: str, ctype: str, raw: bytes) -> str:
        lower = url.lower()
        if "pdf" in ctype or lower.endswith(".pdf") or raw[:4] == b"%PDF":
            return "application/pdf"
        if "json" in ctype or lower.endswith(".json"):
            return "application/json"
        if "csv" in ctype or lower.endswith(".csv"):
            return "text/csv"
        if "sheet" in ctype or lower.endswith(".xlsx") or lower.endswith(".xls"):
            return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if "xml" in ctype or "rss" in ctype or lower.endswith(".xml") or lower.endswith(".rss"):
            return "application/xml"
        if "html" in ctype or lower.endswith(".html") or lower.endswith(".htm"):
            return "text/html"
        # sniff
        head = raw[:200].lstrip().lower()
        if head.startswith(b"<!doctype html") or head.startswith(b"<html"):
            return "text/html"
        if head.startswith(b"{") or head.startswith(b"["):
            return "application/json"
        if head.startswith(b"<?xml") or b"<rss" in head:
            return "application/xml"
        return ctype.split(";")[0].strip() if ctype else "application/octet-stream"

    def _extract_text(self, detected: str, raw: bytes, url: str) -> str:
        if detected == "application/pdf":
            return self._extract_pdf(raw) if self.pdf_enabled else ""
        if detected in {"application/json"}:
            try:
                return json.dumps(json.loads(raw.decode("utf-8", errors="ignore")), indent=2)[:100_000]
            except Exception:
                return raw.decode("utf-8", errors="ignore")
        if detected in {"text/csv"}:
            return raw.decode("utf-8", errors="ignore")
        if detected in {"application/xml", "application/rss+xml"} or "xml" in detected:
            return self._extract_xml_rss(raw)
        if "sheet" in detected or detected.endswith("xlsx"):
            return self._extract_xlsx(raw)
        return self._html_to_text(raw.decode("utf-8", errors="ignore"))

    def _extract_xml_rss(self, raw: bytes) -> str:
        try:
            root = ElementTree.fromstring(raw)
        except Exception:
            return raw.decode("utf-8", errors="ignore")
        lines = [f"XML/RSS root: {root.tag}"]
        for item in list(root.iter())[:80]:
            tag = item.tag.split("}")[-1].lower()
            if tag in {"title", "link", "description", "pubdate", "guid"} and (item.text or "").strip():
                lines.append(f"{tag}: {(item.text or '').strip()}")
        return "\n".join(lines)

    def _extract_xlsx(self, raw: bytes) -> str:
        # Best-effort: avoid hard dependency; store size provenance if openpyxl missing
        try:
            import openpyxl  # type: ignore

            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets[:3]:
                lines.append(f"[sheet {sheet.title}]")
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i > 40:
                        break
                    vals = [str(c) for c in row if c is not None]
                    if vals:
                        lines.append(" | ".join(vals))
            return "\n".join(lines)
        except Exception as exc:
            return f"[xlsx_extract_unavailable] bytes={len(raw)} error={exc}"

    def _fetch_search(self, candidate: CandidateDocument) -> FetchedDocument:
        meta = candidate.metadata or {}
        if meta.get("deferred") or not meta.get("providers_available"):
            return FetchedDocument(
                candidate_id=candidate.candidate_id,
                title=candidate.title,
                url=candidate.url,
                connector_id=candidate.connector_id,
                document_type=candidate.document_type,
                company=candidate.company,
                symbol=candidate.symbol,
                skipped=True,
                skip_reason="search_api_unconfigured",
                live_fetch=False,
                metadata=meta,
            )
        if not self.live_fetch:
            return FetchedDocument(
                candidate_id=candidate.candidate_id,
                title=candidate.title,
                url=candidate.url,
                connector_id=candidate.connector_id,
                document_type=candidate.document_type,
                company=candidate.company,
                symbol=candidate.symbol,
                skipped=True,
                skip_reason="live_fetch_disabled",
                live_fetch=False,
                metadata=meta,
            )

        # Prefer provider encoded in search://{provider}?q=...
        provider = str(meta.get("selected_provider") or "").strip()
        if not provider:
            try:
                from urllib.parse import urlparse

                host = urlparse(candidate.url).netloc or ""
                path_provider = (urlparse(candidate.url).path or "").strip("/")
                # search://exa?q=... → netloc=exa
                provider = host or path_provider.split("?")[0]
            except Exception:
                provider = ""
        if not provider:
            provider = (meta.get("providers_available") or ["exa", "tavily"])[0]
        query = meta.get("query") or candidate.title
        t0 = time.perf_counter()
        results = self._call_search_provider(provider, str(query))
        # Deepen top hits with Firecrawl/Browserbase markdown (strategic, capped)
        results = deepen_search_results(self.client, results, max_pages=3)
        fetch_ms = (time.perf_counter() - t0) * 1000
        self._note_fetch(fetch_ms)
        lines = [f"Search provider: {provider}", f"Query: {query}", ""]
        for i, r in enumerate(results[:8], 1):
            lines.append(f"{i}. {r.get('title')}")
            lines.append(f"   URL: {r.get('url')}")
            if r.get("snippet"):
                lines.append(f"   {r.get('snippet')}")
            if r.get("enriched_by"):
                lines.append(f"   enriched_by: {r.get('enriched_by')}")
            if r.get("markdown"):
                # Include a bounded body so FRE indexes real page content, not only snippets
                body = str(r.get("markdown"))[:6_000].strip()
                if body:
                    lines.append("   --- page markdown ---")
                    lines.append(f"   {body}")
            lines.append("")
        text = "\n".join(lines).strip() or f"No results from {provider} for {query}"
        checksum = sha256_text(text)
        enriched_n = sum(1 for r in results[:8] if r.get("enriched_by") or r.get("markdown"))
        return FetchedDocument(
            candidate_id=candidate.candidate_id,
            title=f"{provider} results — {query}",
            url=candidate.url,
            connector_id=candidate.connector_id,
            document_type=candidate.document_type,
            company=candidate.company,
            symbol=candidate.symbol,
            organisation=provider,
            content_type="text/markdown" if enriched_n else "application/json",
            content_text=text,
            content_bytes_len=len(text.encode("utf-8")),
            checksum=checksum,
            live_fetch=True,
            fetch_ms=fetch_ms,
            metadata={
                "provider": provider,
                "result_count": len(results),
                "results": [
                    {
                        **{k: v for k, v in r.items() if k != "markdown"},
                        **({"has_markdown": True} if r.get("markdown") else {}),
                    }
                    for r in results[:8]
                ],
                "pages_enriched": enriched_n,
            },
        )

    def _call_search_provider(self, provider: str, query: str) -> list[dict[str, Any]]:
        if provider == "tavily":
            key = (os.environ.get("TAVILY_API_KEY") or "").strip()
            if not key:
                return []
            resp = self.client.post_json(
                "https://api.tavily.com/search",
                {"api_key": key, "query": query, "max_results": 5, "include_answer": False},
                connector_id="tavily",
            )
            if resp.error or not resp.ok:
                return []
            data = json.loads(resp.text or "{}")
            return [
                {"title": r.get("title"), "url": r.get("url"), "snippet": r.get("content")}
                for r in (data.get("results") or [])
                if isinstance(r, dict)
            ]

        if provider == "serpapi":
            key = (os.environ.get("SERPAPI_API_KEY") or os.environ.get("SERPAPI_KEY") or "").strip()
            if not key:
                return []
            resp = self.client.get(
                f"https://serpapi.com/search.json?q={query}&api_key={key}&engine=google&num=5",
                connector_id="serpapi",
                max_per_minute=20,
            )
            if resp.error or not resp.ok:
                return []
            data = json.loads(resp.text or "{}")
            return [
                {"title": r.get("title"), "url": r.get("link"), "snippet": r.get("snippet")}
                for r in (data.get("organic_results") or [])
                if isinstance(r, dict)
            ]

        if provider == "exa":
            from app.faa.provider_flags import provider_enabled

            if not provider_enabled("exa"):
                return []
            key = (os.environ.get("EXA_API_KEY") or "").strip()
            if not key:
                return []
            resp = self.client.post_json(
                "https://api.exa.ai/search",
                {
                    "query": query,
                    "numResults": 6,
                    "type": "auto",
                    "contents": {"text": {"maxCharacters": 1200}},
                    "useAutoprompt": True,
                },
                connector_id="exa",
                headers={"x-api-key": key},
            )
            if resp.error or not resp.ok:
                return []
            data = json.loads(resp.text or "{}")
            out: list[dict[str, Any]] = []
            for r in data.get("results") or []:
                if not isinstance(r, dict):
                    continue
                snippet = (r.get("text") or "").strip()
                if not snippet and isinstance(r.get("highlights"), list) and r["highlights"]:
                    snippet = str(r["highlights"][0])
                out.append(
                    {
                        "title": r.get("title"),
                        "url": r.get("url"),
                        "snippet": snippet[:500],
                    }
                )
            return out

        if provider == "firecrawl":
            from app.faa.web_enrichment import firecrawl_search

            return firecrawl_search(self.client, query, limit=5)

        if provider == "playwright":
            from app.faa.playwright_browser import web_search as playwright_web_search

            return playwright_web_search(query, limit=5)

        if provider == "bing":
            key = (os.environ.get("BING_SEARCH_API_KEY") or "").strip()
            if not key:
                return []
            resp = self.client.get(
                f"https://api.bing.microsoft.com/v7.0/search?q={query}&count=5",
                connector_id="bing",
                headers={"Ocp-Apim-Subscription-Key": key},
                max_per_minute=20,
            )
            if resp.error or not resp.ok:
                return []
            data = json.loads(resp.text or "{}")
            return [
                {"title": r.get("name"), "url": r.get("url"), "snippet": r.get("snippet")}
                for r in ((data.get("webPages") or {}).get("value") or [])
                if isinstance(r, dict)
            ]

        if provider == "google_cse":
            key = (os.environ.get("GOOGLE_CSE_API_KEY") or "").strip()
            cx = (os.environ.get("GOOGLE_CSE_ID") or "").strip()
            if not key or not cx:
                return []
            resp = self.client.get(
                f"https://www.googleapis.com/customsearch/v1?key={key}&cx={cx}&q={query}&num=5",
                connector_id="google_cse",
                max_per_minute=20,
            )
            if resp.error or not resp.ok:
                return []
            data = json.loads(resp.text or "{}")
            return [
                {"title": r.get("title"), "url": r.get("link"), "snippet": r.get("snippet")}
                for r in (data.get("items") or [])
                if isinstance(r, dict)
            ]
        return []

    def _extract_pdf(self, raw: bytes) -> str:
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(raw))
            pages = []
            for i, page in enumerate(reader.pages[:50]):
                try:
                    t = page.extract_text() or ""
                except Exception:
                    t = ""
                if t.strip():
                    pages.append(f"[page {i+1}]\n{t.strip()}")
            return "\n\n".join(pages)
        except Exception as exc:
            return f"[pdf_extract_failed] {exc}"

    def _html_to_text(self, html: str) -> str:
        html = re.sub(r"(?is)<(script|style|noscript).*?>.*?</\1>", " ", html)
        html = re.sub(r"(?is)<!--.*?-->", " ", html)
        text = _TAG_RE.sub(" ", html)
        text = _WS_RE.sub(" ", text)
        return text.strip()

    def _offline_stub_text(self, candidate: CandidateDocument) -> str:
        return (
            f"{candidate.title}\n"
            f"Source URL: {candidate.url}\n"
            f"Connector: {candidate.connector_id}\n"
            f"Document type: {candidate.document_type}\n"
            f"Company: {candidate.company or 'n/a'} ({candidate.symbol or 'n/a'})\n"
            f"Organisation: {candidate.organisation or 'n/a'}\n\n"
            "FAA offline acquisition record. Enable FAA_LIVE_FETCH=true to download the live page/PDF. "
            "This preserves the discovery target so FRE can track intended public sources."
        )

    def _fail(
        self,
        candidate: CandidateDocument,
        error: str,
        *,
        fetch_ms: float = 0.0,
        attempts: int = 1,
    ) -> FetchedDocument:
        return FetchedDocument(
            candidate_id=candidate.candidate_id,
            title=candidate.title,
            url=candidate.url,
            connector_id=candidate.connector_id,
            document_type=candidate.document_type,
            company=candidate.company,
            symbol=candidate.symbol,
            organisation=candidate.organisation,
            error=error,
            live_fetch=self.live_fetch,
            fetch_ms=fetch_ms,
            attempts=attempts,
        )

    def _note_fetch(self, ms: float) -> None:
        self._fetch_samples.append(ms)
        self._fetch_samples = self._fetch_samples[-200:]

    @property
    def avg_fetch_ms(self) -> float:
        if not self._fetch_samples:
            return 0.0
        return sum(self._fetch_samples) / len(self._fetch_samples)
